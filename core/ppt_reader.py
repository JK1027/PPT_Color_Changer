import os
import logging
from typing import Set
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# 모듈화된 서브시스템 임포트
from core.ppt_scanner import extract_all_colors
from core.text_color_replacer import replace_text_color_in_shape
from core.image_replacer import replace_image_color_in_shape

logger = logging.getLogger("ppt_color_changer")

def replace_color(
    ppt_path: str, 
    target_color: str, 
    new_color_hex: str, 
    output_path: str, 
    change_text: bool = True, 
    change_image: bool = True,
    tolerance: int = 20
) -> int:
    """
    특정 색상을 새 색상으로 일괄 변경하고 저장 
    (텍스트 및 이미지 색상 치환 선택적 포함, 오차 허용치 설정 가능)
    (원본 보존 원칙: 무조건 새로운 output_path로 저장)
    """
    if os.path.abspath(ppt_path) == os.path.abspath(output_path):
        raise ValueError("원본 데이터 보존을 위해 output_path는 ppt_path와 동일할 수 없습니다.")

    try:
        prs = Presentation(ppt_path)
        change_count = 0
        
        for slide in prs.slides:
            # slide.shapes를 리스트로 복사하여 순회해야 shape 삭제/추가 시 인덱스 에러가 안남
            for shape in list(slide.shapes):
                change_count += _replace_color_in_shape(
                    shape, slide, target_color, new_color_hex, change_text, change_image, tolerance
                )
                
        prs.save(output_path)
        logger.info(f"Successfully replaced colors in PPT: {ppt_path} -> {output_path} (Changes: {change_count})")
        return change_count
    except Exception as e:
        logger.error(f"Failed to process presentation {ppt_path}: {e}", exc_info=True)
        raise

def _replace_color_in_shape(
    shape, 
    slide, 
    target_color: str, 
    new_color_hex: str, 
    change_text: bool, 
    change_image: bool,
    tolerance: int
) -> int:
    """재귀적으로 shape 탐색하여 색상 변경 후 변경 횟수 반환"""
    change_count = 0
    
    # 1. 텍스트 변경 (일반 도형 또는 표)
    if change_text and (shape.has_text_frame or shape.has_table):
        change_count += replace_text_color_in_shape(shape, target_color, new_color_hex)
        
    # 2. 이미지 변경 (PICTURE)
    elif change_image and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        change_count += replace_image_color_in_shape(shape, slide, target_color, new_color_hex, tolerance)
        
    # 3. 그룹 도형인 경우 재귀 탐색 (텍스트/이미지 옵션을 하위 도형에 그대로 위임)
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child_shape in shape.shapes:
            change_count += _replace_color_in_shape(
                child_shape, slide, target_color, new_color_hex, change_text, change_image, tolerance
            )
            
    return change_count
