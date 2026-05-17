import io
import logging
from pptx.enum.shapes import MSO_SHAPE_TYPE
from core.image_processor import process_image_blob

logger = logging.getLogger("ppt_color_changer")

def replace_image_color_in_shape(shape, slide, target_color: str, new_color_hex: str, tolerance: int = 20) -> int:
    """
    [실험적 기능 (Experimental Image Swap)]
    이미지 내의 특정 색상을 HSV 기반으로 치환하고 슬라이드의 shape를 스왑합니다.
    
    NOTE: python-pptx는 자체적인 이미지 '교체' API가 빈약하기 때문에, 동일 위치/크기로
    새 이미지를 추가한 뒤 lxml 트리 레벨에서 기존 shape 엘리먼트를 강제 제거하는 기법을 사용합니다.
    이 로직은 python-pptx의 내부 구현에 의존하므로 예외 상황 발생 시 원래 이미지를 안전하게 
    유지하도록 try-except로 완전히 감싸 격리시킵니다.
    """
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return 0
        
    try:
        image_blob = shape.image.blob
        
        # 이미지 프로세서를 거쳐 색상 치환된 바이너리 획득 (유사도 오차 반영)
        new_blob = process_image_blob(image_blob, target_color, new_color_hex, tolerance)
        
        # 바이너리가 달라졌다면 (실제 색상 치환이 일어났다면) 이미지 교체
        if new_blob != image_blob:
            image_stream = io.BytesIO(new_blob)
            # 동일한 위치와 크기(left, top, width, height)로 새 이미지 추가
            slide.shapes.add_picture(image_stream, shape.left, shape.top, shape.width, shape.height)
            
            # lxml 엘리먼트 획득 및 부모 노드에서 원본 제거 (Workaround)
            sp = shape._element
            sp.getparent().remove(sp)
            logger.info("Successfully swapped image shape via experimental XML DOM manipulation.")
            return 1
            
    except Exception as e:
        logger.error(f"[Experimental Swap Warning] Failed to replace picture shape: {e}", exc_info=True)
        
    return 0
